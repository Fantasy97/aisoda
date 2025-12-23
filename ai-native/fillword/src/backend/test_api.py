#!/usr/bin/env python3
"""
API测试脚本 - 在 src/backend 目录下运行
"""

import requests
import json
import os
import time

BASE_URL = "http://localhost:5000"

def test_health_check():
    """测试健康检查"""
    print("🔍 测试健康检查...")
    try:
        response = requests.get(f"{BASE_URL}/api/health")
        if response.status_code == 200:
            result = response.json()
            print(f"✅ 健康检查通过: {result['status']}")
            return True
        else:
            print(f"❌ 健康检查失败: {response.status_code}")
            return False
    except Exception as e:
        print(f"❌ 健康检查异常: {e}")
        return False

def test_single_query():
    """测试单个查询"""
    print("\n🔍 测试单个查询...")
    try:
        data = {"keyword": "企业名称"}
        response = requests.post(
            f"{BASE_URL}/api/query",
            json=data,
            headers={"Content-Type": "application/json"}
        )
        
        if response.status_code == 200:
            result = response.json()
            if result['success']:
                print(f"✅ 查询成功: {result['keyword']} -> {result['result']}")
                return True
            else:
                print(f"❌ 查询失败: {result['error']}")
                return False
        else:
            print(f"❌ 查询请求失败: {response.status_code}")
            return False
    except Exception as e:
        print(f"❌ 查询异常: {e}")
        return False

def create_test_document():
    """创建测试文档"""
    print("\n📝 创建测试文档...")
    try:
        from docx import Document
        
        doc = Document()
        doc.add_heading('测试表单', 0)
        
        table = doc.add_table(rows=3, cols=2)
        table.style = 'Table Grid'
        
        table.cell(0, 0).text = '企业名称'
        table.cell(0, 1).text = ''
        table.cell(1, 0).text = '联系电话'
        table.cell(1, 1).text = ''
        table.cell(2, 0).text = '法定代表人'
        table.cell(2, 1).text = ''
        
        doc_path = "test_document.docx"
        doc.save(doc_path)
        print(f"✅ 测试文档创建成功: {doc_path}")
        return doc_path
        
    except ImportError:
        print("❌ 无法创建测试文档: 缺少python-docx库")
        return None
    except Exception as e:
        print(f"❌ 创建测试文档失败: {e}")
        return None

def test_document_upload(doc_path):
    """测试文档上传"""
    print(f"\n🔍 测试文档上传: {doc_path}")
    
    if not os.path.exists(doc_path):
        print(f"❌ 测试文档不存在: {doc_path}")
        return False
    
    try:
        with open(doc_path, 'rb') as f:
            files = {'file': f}
            print("⏳ 正在上传和处理文档...")
            response = requests.post(f"{BASE_URL}/api/upload", files=files)
        
        if response.status_code == 200:
            result = response.json()
            if result['success']:
                stats = result['statistics']
                print(f"✅ 文档处理成功!")
                print(f"   - 总处理数量: {stats['total_processed']}")
                print(f"   - 成功填充: {stats['successful_fills']}")
                print(f"   - 成功率: {stats['success_rate']:.1%}")
                
                # 测试下载
                download_url = f"{BASE_URL}{result['download_url']}"
                download_response = requests.get(download_url)
                
                if download_response.status_code == 200:
                    # 确保输出目录存在
                    output_dir = "../../data/output"
                    os.makedirs(output_dir, exist_ok=True)
                    output_path = os.path.join(output_dir, f"downloaded_{result['output_filename']}")
                    with open(output_path, 'wb') as f:
                        f.write(download_response.content)
                    print(f"✅ 文件下载成功: {output_path}")
                    
                    # 检查输出目录
                    output_dir = "../../data/output"
                    processed_dir = "../../data/processed"
                    if os.path.exists(output_dir):
                        output_files = os.listdir(output_dir)
                        print(f"📁 输出目录文件: {output_files}")
                    if os.path.exists(processed_dir):
                        processed_files = os.listdir(processed_dir)
                        print(f"📁 处理目录文件: {processed_files}")
                
                return True
            else:
                print(f"❌ 文档处理失败: {result['error']}")
                return False
        else:
            print(f"❌ 上传请求失败: {response.status_code}")
            return False
            
    except Exception as e:
        print(f"❌ 上传异常: {e}")
        return False

def main():
    """主测试函数"""
    print("🚀 开始API测试")
    print("=" * 40)
    
    # 检查服务状态
    if not test_health_check():
        print("\n❌ 服务未启动，请先运行: python run.py")
        return
    
    time.sleep(1)
    
    # 测试单个查询
    test_single_query()
    
    # 创建并测试文档上传
    test_doc_path = create_test_document()
    if test_doc_path:
        test_document_upload(test_doc_path)
        
        # 清理测试文件
        try:
            os.remove(test_doc_path)
        except:
            pass
    
    print("\n🎉 测试完成！")

if __name__ == '__main__':
    main()