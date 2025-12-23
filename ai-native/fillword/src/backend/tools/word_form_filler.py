"""
Word表单填充工具
将JSON数据填写到Word文档中，并生成填充结果报告
采用模块化设计，包含文档处理器、填充引擎和报告生成器
"""

import json
import logging
import argparse
from pathlib import Path
from typing import Dict, List, Any, Optional
from dataclasses import dataclass
from datetime import datetime
import os

try:
    from docx import Document
    from docx.table import Table, _Cell
    from docx.shared import Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH
except ImportError:
    print("请安装python-docx库: pip install python-docx")
    exit(1)


@dataclass
class FillResult:
    """填充结果数据结构"""
    row: int
    col: int
    original_content: str
    new_content: str
    relationship: List[str]
    fill_source: str
    confidence: float
    success: bool
    error_message: str = ""


class DocumentProcessor:
    """文档处理器 - 负责Word文档的读取、单元格定位和内容填充"""
    
    def __init__(self):
        self.logger = logging.getLogger(f"{__name__}.DocumentProcessor")
    
    def load_document(self, word_path: str) -> Optional[Document]:
        """加载Word文档"""
        try:
            if not os.path.exists(word_path):
                self.logger.error(f"Word文档不存在: {word_path}")
                return None
            
            doc = Document(word_path)
            self.logger.info(f"成功加载Word文档: {word_path}")
            return doc
            
        except Exception as e:
            self.logger.error(f"加载Word文档失败: {e}")
            return None
    
    def find_table_cell(self, doc: Document, row: int, col: int) -> Optional[_Cell]:
        """在文档中查找指定位置的单元格"""
        try:
            # 遍历文档中的所有表格
            for table_idx, table in enumerate(doc.tables):
                if row < len(table.rows) and col < len(table.rows[row].cells):
                    cell = table.rows[row].cells[col]
                    self.logger.debug(f"找到单元格 表格{table_idx} 行{row} 列{col}")
                    return cell
            
            self.logger.warning(f"未找到单元格: 行{row} 列{col}")
            return None
            
        except Exception as e:
            self.logger.error(f"查找单元格失败 行{row} 列{col}: {e}")
            return None
    
    def fill_cell_content(self, cell: _Cell, content: str) -> bool:
        """填充单元格内容"""
        try:
            if not cell:
                return False
            
            # 清空原有内容
            cell.text = ""
            
            # 添加新内容
            paragraph = cell.paragraphs[0]
            paragraph.text = content
            
            # 设置对齐方式
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            
            return True
            
        except Exception as e:
            self.logger.error(f"填充单元格失败: {e}")
            return False
    
    def save_document(self, doc: Document, output_path: str) -> bool:
        """保存文档"""
        try:
            doc.save(output_path)
            self.logger.info(f"文档保存成功: {output_path}")
            return True
        except Exception as e:
            self.logger.error(f"保存文档失败: {e}")
            return False


class DataLoader:
    """数据加载器 - 负责JSON数据的加载和验证"""
    
    def __init__(self):
        self.logger = logging.getLogger(f"{__name__}.DataLoader")
    
    def load_json_data(self, json_path: str) -> List[Dict[str, Any]]:
        """加载JSON数据"""
        try:
            if not os.path.exists(json_path):
                self.logger.error(f"JSON文件不存在: {json_path}")
                return []
            
            with open(json_path, 'r', encoding='utf-8') as f:
                data = json.load(f)
            
            if not isinstance(data, list):
                self.logger.error("JSON数据必须是数组格式")
                return []
            
            self.logger.info(f"成功加载JSON数据: {len(data)} 条记录")
            return data
            
        except json.JSONDecodeError as e:
            self.logger.error(f"JSON格式错误: {e}")
            return []
        except Exception as e:
            self.logger.error(f"加载JSON数据失败: {e}")
            return []
    
    def validate_data_item(self, item: Dict[str, Any]) -> bool:
        """验证单个数据项的格式"""
        required_fields = ['row', 'col', 'content']
        
        for field in required_fields:
            if field not in item:
                self.logger.warning(f"数据项缺少必需字段: {field}")
                return False
        
        if not isinstance(item['row'], int) or item['row'] < 0:
            self.logger.warning(f"无效的行号: {item['row']}")
            return False
        
        if not isinstance(item['col'], int) or item['col'] < 0:
            self.logger.warning(f"无效的列号: {item['col']}")
            return False
        
        return True


class FillEngine:
    """填充引擎 - 协调文档处理器和数据加载器完成填充任务"""
    
    def __init__(self, log_level: str = "INFO"):
        """初始化填充引擎"""
        self.setup_logging(log_level)
        self.logger = logging.getLogger(f"{__name__}.FillEngine")
        
        # 初始化组件
        self.document_processor = DocumentProcessor()
        self.data_loader = DataLoader()
        self.fill_results: List[FillResult] = []
        
    def setup_logging(self, level: str) -> None:
        """设置日志"""
        # 确保日志目录存在
        log_dir = '../../data/processed'
        os.makedirs(log_dir, exist_ok=True)
        
        logging.basicConfig(
            level=getattr(logging, level.upper()),
            format='%(asctime)s - %(name)s - %(levelname)s - %(message)s',
            handlers=[
                logging.FileHandler(os.path.join(log_dir, 'word_form_filler.log'), encoding='utf-8'),
                logging.StreamHandler()
            ]
        )
    
    def load_json_data(self, json_path: str) -> List[Dict[str, Any]]:
        """加载JSON数据（兼容性方法）"""
        return self.data_loader.load_json_data(json_path)
    
    def fill_word_document(self, word_path: str, json_data: List[Dict[str, Any]], output_path: str = None) -> bool:
        """填充Word文档的主要方法"""
        try:
            # 加载Word文档
            doc = self.document_processor.load_document(word_path)
            if not doc:
                return False
            
            # 如果没有指定输出路径，生成默认路径
            if not output_path:
                word_file = Path(word_path)
                output_path = str(word_file.parent / f"{word_file.stem}_filled{word_file.suffix}")
            
            # 执行填充处理
            success = self._process_fill_data(doc, json_data)
            
            if success:
                # 保存填充后的文档
                return self.document_processor.save_document(doc, output_path)
            
            return False
            
        except Exception as e:
            self.logger.error(f"填充Word文档失败: {e}")
            return False
    
    def _process_fill_data(self, doc: Document, json_data: List[Dict[str, Any]]) -> bool:
        """处理填充数据的内部方法"""
        try:
            # 统计信息
            total_items = len(json_data)
            filled_count = 0
            error_count = 0
            
            # 遍历JSON数据进行填充
            for item in json_data:
                try:
                    # 验证数据项
                    if not self.data_loader.validate_data_item(item):
                        error_count += 1
                        continue
                    
                    # 提取数据
                    row = item.get('row', 0)
                    col = item.get('col', 0)
                    content = item.get('content', '')
                    relationship = item.get('relationship', [])
                    fill_source = item.get('fill_source', '')
                    confidence = item.get('confidence', 0.0)
                    
                    # 跳过空内容
                    if not content.strip():
                        self.logger.debug(f"跳过空内容: 行{row} 列{col}")
                        continue
                    
                    # 执行单个单元格填充
                    result = self._fill_single_cell(doc, item)
                    self.fill_results.append(result)
                    
                    if result.success:
                        filled_count += 1
                        self.logger.info(f"填充成功: 行{row} 列{col} -> {content}")
                    else:
                        error_count += 1
                        self.logger.error(f"填充失败: 行{row} 列{col} - {result.error_message}")
                
                except Exception as e:
                    error_count += 1
                    self.logger.error(f"处理数据项失败: {e}")
                    continue
            
            # 打印统计信息
            self.logger.info(f"填充完成 - 总数: {total_items}, 成功: {filled_count}, 失败: {error_count}")
            
            return filled_count > 0
            
        except Exception as e:
            self.logger.error(f"处理填充数据失败: {e}")
            return False
    
    def _fill_single_cell(self, doc: Document, item: Dict[str, Any]) -> FillResult:
        """填充单个单元格"""
        row = item.get('row', 0)
        col = item.get('col', 0)
        content = item.get('content', '')
        relationship = item.get('relationship', [])
        fill_source = item.get('fill_source', '')
        confidence = item.get('confidence', 0.0)
        
        try:
            # 查找目标单元格
            cell = self.document_processor.find_table_cell(doc, row, col)
            
            if cell:
                # 获取原始内容
                original_content = cell.text
                
                # 填充内容
                success = self.document_processor.fill_cell_content(cell, content)
                
                return FillResult(
                    row=row,
                    col=col,
                    original_content=original_content,
                    new_content=content,
                    relationship=relationship,
                    fill_source=fill_source,
                    confidence=confidence,
                    success=success,
                    error_message="" if success else "填充操作失败"
                )
            
            else:
                return FillResult(
                    row=row,
                    col=col,
                    original_content="",
                    new_content=content,
                    relationship=relationship,
                    fill_source=fill_source,
                    confidence=confidence,
                    success=False,
                    error_message="未找到目标单元格"
                )
        
        except Exception as e:
            return FillResult(
                row=row,
                col=col,
                original_content="",
                new_content=content,
                relationship=relationship,
                fill_source=fill_source,
                confidence=confidence,
                success=False,
                error_message=str(e)
            )
    
    def get_statistics(self) -> Dict[str, Any]:
        """获取统计信息"""
        total_count = len(self.fill_results)
        success_count = sum(1 for r in self.fill_results if r.success)
        error_count = total_count - success_count
        
        # 按填充源分组统计
        source_stats = {}
        for result in self.fill_results:
            source = result.fill_source.split(':')[0] if ':' in result.fill_source else result.fill_source
            if source not in source_stats:
                source_stats[source] = {'total': 0, 'success': 0}
            source_stats[source]['total'] += 1
            if result.success:
                source_stats[source]['success'] += 1
        
        # 置信度统计
        confidences = [r.confidence for r in self.fill_results if r.confidence > 0]
        avg_confidence = sum(confidences) / len(confidences) if confidences else 0
        
        return {
            'total_processed': total_count,
            'successful_fills': success_count,
            'failed_fills': error_count,
            'success_rate': (success_count / total_count) if total_count > 0 else 0,
            'source_statistics': source_stats,
            'average_confidence': avg_confidence,
            'confidence_distribution': {
                'high (>0.9)': len([c for c in confidences if c > 0.9]),
                'medium (0.7-0.9)': len([c for c in confidences if 0.7 <= c <= 0.9]),
                'low (<0.7)': len([c for c in confidences if c < 0.7])
            }
        }


class ReportGenerator:
    """报告生成器 - 负责生成HTML格式的填充结果报告"""
    
    def __init__(self):
        self.logger = logging.getLogger(f"{__name__}.ReportGenerator")
    
    def generate_report(self, fill_results: List[FillResult], report_path: str = None, download_url: str = None) -> bool:
        """生成HTML格式的填充结果报告"""
        try:
            if not fill_results:
                self.logger.warning("没有填充结果数据，无法生成报告")
                return False
            
            if not report_path:
                timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
                # 输出到output目录
                output_dir = Path("../../data/output")
                output_dir.mkdir(parents=True, exist_ok=True)
                report_path = str(output_dir / f"填充报告_{timestamp}.html")
            
            # 确保文件扩展名为.html
            if not report_path.lower().endswith('.html'):
                report_path = os.path.splitext(report_path)[0] + '.html'
            
            # 计算统计信息
            total_count = len(fill_results)
            success_count = sum(1 for r in fill_results if r.success)
            error_count = total_count - success_count
            success_rate = (success_count / total_count * 100) if total_count > 0 else 0
            
            # 生成HTML内容
            html_content = self._generate_html_content(fill_results, total_count, success_count, error_count, success_rate, download_url)
            
            # 保存HTML文件
            with open(report_path, 'w', encoding='utf-8') as f:
                f.write(html_content)
            
            self.logger.info(f"HTML报告生成成功: {report_path}")
            return True
            
        except Exception as e:
            self.logger.error(f"生成HTML报告失败: {e}")
            return False
    
    def _generate_html_content(self, fill_results: List[FillResult], total_count: int, 
                              success_count: int, error_count: int, success_rate: float, download_url: str = None) -> str:
        """生成HTML内容"""
        current_time = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        
        # 按填充源分组统计
        source_stats = self._calculate_source_statistics(fill_results)
        
        # 置信度统计
        confidence_stats = self._calculate_confidence_statistics(fill_results)
        
        html_content = f"""
<!DOCTYPE html>
<html lang="zh-CN">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>Word表单填充结果报告</title>
    <style>
        {self._get_css_styles()}
    </style>
</head>
<body>
    <div class="container">
        <header>
            <h1>📋 Word表单填充结果报告</h1>
            <p class="timestamp">报告生成时间: {current_time}</p>
        </header>
        
        <section class="summary">
            <h2>📊 统计摘要</h2>
            <div class="stats-grid">
                <div class="stat-card">
                    <div class="stat-number">{total_count}</div>
                    <div class="stat-label">总处理数量</div>
                </div>
                <div class="stat-card success">
                    <div class="stat-number">{success_count}</div>
                    <div class="stat-label">成功填充</div>
                </div>
                <div class="stat-card error">
                    <div class="stat-number">{error_count}</div>
                    <div class="stat-label">失败数量</div>
                </div>
                <div class="stat-card rate">
                    <div class="stat-number">{success_rate:.1f}%</div>
                    <div class="stat-label">成功率</div>
                </div>
            </div>
        </section>
        
        <section class="source-stats">
            <h2>📈 按填充源统计</h2>
            {self._generate_source_stats_html(source_stats)}
        </section>
        
        <section class="confidence-stats">
            <h2>🎯 置信度分布</h2>
            {self._generate_confidence_stats_html(confidence_stats)}
        </section>
        
        <section class="detailed-results">
            <h2>📋 详细结果</h2>
            {self._generate_detailed_results_html(fill_results)}
        </section>
        
        {self._generate_error_details_html(fill_results)}
        
        <footer>
            <p>报告由 Word表单填充工具 自动生成</p>
        </footer>
    </div>
    
    <!-- 导航按钮 -->
    <div class="navigation-buttons" style="position: fixed; top: 20px; right: 20px; z-index: 1000; display: flex; gap: 10px;">
        <a href="/" class="btn btn-secondary" style="background: #6c757d; color: white; padding: 10px 15px; text-decoration: none; border-radius: 4px; font-size: 14px;">
            🏠 返回主页
        </a>
        {f'<a href="{download_url}" class="btn btn-primary" style="background: #007bff; color: white; padding: 10px 15px; text-decoration: none; border-radius: 4px; font-size: 14px;">📥 下载文档</a>' if download_url else ''}
    </div>
    
    <script>
        // 报告页面增强脚本
        document.addEventListener('DOMContentLoaded', function() {{
            // 添加下载按钮样式
            const style = document.createElement('style');
            style.textContent = `
                .btn {{
                    display: inline-block;
                    padding: 10px 20px;
                    margin: 5px;
                    text-decoration: none;
                    border-radius: 4px;
                    font-size: 14px;
                    cursor: pointer;
                    border: none;
                    transition: all 0.3s ease;
                }}
                .btn-primary {{
                    background-color: #007bff;
                    color: white;
                }}
                .btn-secondary {{
                    background-color: #6c757d;
                    color: white;
                }}
                .btn:hover {{
                    opacity: 0.8;
                    transform: translateY(-1px);
                }}
                @media (max-width: 768px) {{
                    .navigation-buttons {{
                        position: static !important;
                        justify-content: center;
                        margin: 20px 0;
                    }}
                }}
            `;
            document.head.appendChild(style);
        }});
    </script>
</body>
</html>
"""
        return html_content
    
    def _get_css_styles(self) -> str:
        """获取CSS样式"""
        return """
        * {
            margin: 0;
            padding: 0;
            box-sizing: border-box;
        }
        
        body {
            font-family: 'Segoe UI', Tahoma, Geneva, Verdana, sans-serif;
            line-height: 1.6;
            color: #333;
            background-color: #f5f5f5;
        }
        
        .container {
            max-width: 1200px;
            margin: 0 auto;
            padding: 20px;
            background-color: white;
            box-shadow: 0 0 10px rgba(0,0,0,0.1);
        }
        
        header {
            text-align: center;
            margin-bottom: 30px;
            padding-bottom: 20px;
            border-bottom: 2px solid #e0e0e0;
        }
        
        h1 {
            color: #2c3e50;
            margin-bottom: 10px;
            font-size: 2.5em;
        }
        
        h2 {
            color: #34495e;
            margin: 30px 0 15px 0;
            font-size: 1.8em;
            border-left: 4px solid #3498db;
            padding-left: 15px;
        }
        
        .timestamp {
            color: #7f8c8d;
            font-size: 1.1em;
        }
        
        .stats-grid {
            display: grid;
            grid-template-columns: repeat(auto-fit, minmax(200px, 1fr));
            gap: 20px;
            margin: 20px 0;
        }
        
        .stat-card {
            background: linear-gradient(135deg, #667eea 0%, #764ba2 100%);
            color: white;
            padding: 25px;
            border-radius: 10px;
            text-align: center;
            box-shadow: 0 4px 6px rgba(0,0,0,0.1);
            transition: transform 0.3s ease;
        }
        
        .stat-card:hover {
            transform: translateY(-5px);
        }
        
        .stat-card.success {
            background: linear-gradient(135deg, #11998e 0%, #38ef7d 100%);
        }
        
        .stat-card.error {
            background: linear-gradient(135deg, #ff416c 0%, #ff4b2b 100%);
        }
        
        .stat-card.rate {
            background: linear-gradient(135deg, #f093fb 0%, #f5576c 100%);
        }
        
        .stat-number {
            font-size: 2.5em;
            font-weight: bold;
            margin-bottom: 5px;
        }
        
        .stat-label {
            font-size: 1.1em;
            opacity: 0.9;
        }
        
        .source-chart, .confidence-chart {
            background: white;
            border-radius: 8px;
            padding: 20px;
            margin: 15px 0;
            box-shadow: 0 2px 4px rgba(0,0,0,0.1);
        }
        
        .chart-item {
            display: flex;
            align-items: center;
            margin: 10px 0;
            padding: 10px;
            background: #f8f9fa;
            border-radius: 5px;
        }
        
        .chart-label {
            min-width: 120px;
            font-weight: bold;
            color: #2c3e50;
        }
        
        .chart-bar {
            flex: 1;
            height: 20px;
            background: #ecf0f1;
            border-radius: 10px;
            margin: 0 15px;
            overflow: hidden;
        }
        
        .chart-fill {
            height: 100%;
            background: linear-gradient(90deg, #3498db, #2ecc71);
            border-radius: 10px;
            transition: width 0.5s ease;
        }
        
        .chart-value {
            min-width: 80px;
            text-align: right;
            font-weight: bold;
            color: #27ae60;
        }
        
        .results-table {
            width: 100%;
            border-collapse: collapse;
            margin: 20px 0;
            background: white;
            border-radius: 8px;
            overflow: hidden;
            box-shadow: 0 2px 4px rgba(0,0,0,0.1);
        }
        
        .results-table th {
            background: linear-gradient(135deg, #667eea 0%, #764ba2 100%);
            color: white;
            padding: 15px 10px;
            text-align: left;
            font-weight: bold;
        }
        
        .results-table td {
            padding: 12px 10px;
            border-bottom: 1px solid #ecf0f1;
        }
        
        .results-table tr:nth-child(even) {
            background-color: #f8f9fa;
        }
        
        .results-table tr:hover {
            background-color: #e3f2fd;
        }
        
        .status-success {
            color: #27ae60;
            font-weight: bold;
        }
        
        .status-error {
            color: #e74c3c;
            font-weight: bold;
        }
        
        .confidence-high {
            color: #27ae60;
            font-weight: bold;
        }
        
        .confidence-medium {
            color: #f39c12;
            font-weight: bold;
        }
        
        .confidence-low {
            color: #e74c3c;
            font-weight: bold;
        }
        
        .error-section {
            background: #fff5f5;
            border: 1px solid #fed7d7;
            border-radius: 8px;
            padding: 20px;
            margin: 20px 0;
        }
        
        .error-item {
            background: white;
            border-left: 4px solid #e74c3c;
            padding: 15px;
            margin: 10px 0;
            border-radius: 0 5px 5px 0;
        }
        
        footer {
            text-align: center;
            margin-top: 40px;
            padding-top: 20px;
            border-top: 1px solid #e0e0e0;
            color: #7f8c8d;
        }
        
        .no-data {
            text-align: center;
            color: #7f8c8d;
            font-style: italic;
            padding: 40px;
        }
        
        @media (max-width: 768px) {
            .container {
                padding: 10px;
            }
            
            .stats-grid {
                grid-template-columns: 1fr;
            }
            
            .results-table {
                font-size: 0.9em;
            }
            
            .results-table th,
            .results-table td {
                padding: 8px 5px;
            }
        }
        """
    
    def _calculate_source_statistics(self, fill_results: List[FillResult]) -> Dict[str, Dict[str, int]]:
        """计算按填充源分组的统计信息"""
        source_stats = {}
        for result in fill_results:
            source = result.fill_source.split(':')[0] if ':' in result.fill_source else result.fill_source
            if not source:
                source = 'unknown'
            
            if source not in source_stats:
                source_stats[source] = {'total': 0, 'success': 0}
            source_stats[source]['total'] += 1
            if result.success:
                source_stats[source]['success'] += 1
        
        return source_stats
    
    def _calculate_confidence_statistics(self, fill_results: List[FillResult]) -> Dict[str, int]:
        """计算置信度分布统计"""
        confidences = [r.confidence for r in fill_results if r.confidence > 0]
        return {
            'high': len([c for c in confidences if c > 0.9]),
            'medium': len([c for c in confidences if 0.7 <= c <= 0.9]),
            'low': len([c for c in confidences if c < 0.7])
        }
    
    def _generate_source_stats_html(self, source_stats: Dict[str, Dict[str, int]]) -> str:
        """生成填充源统计的HTML"""
        if not source_stats:
            return '<div class="no-data">暂无填充源统计数据</div>'
        
        html = '<div class="source-chart">'
        for source, data in source_stats.items():
            total = data['total']
            success = data['success']
            rate = (success / total * 100) if total > 0 else 0
            
            html += f'''
            <div class="chart-item">
                <div class="chart-label">{source}</div>
                <div class="chart-bar">
                    <div class="chart-fill" style="width: {rate}%"></div>
                </div>
                <div class="chart-value">{success}/{total} ({rate:.1f}%)</div>
            </div>
            '''
        html += '</div>'
        return html
    
    def _generate_confidence_stats_html(self, confidence_stats: Dict[str, int]) -> str:
        """生成置信度统计的HTML"""
        total_confidence = sum(confidence_stats.values())
        if total_confidence == 0:
            return '<div class="no-data">暂无置信度统计数据</div>'
        
        html = '<div class="confidence-chart">'
        
        confidence_levels = [
            ('high', '高 (>0.9)', '#27ae60'),
            ('medium', '中 (0.7-0.9)', '#f39c12'),
            ('low', '低 (<0.7)', '#e74c3c')
        ]
        
        for level, label, color in confidence_levels:
            count = confidence_stats.get(level, 0)
            rate = (count / total_confidence * 100) if total_confidence > 0 else 0
            
            html += f'''
            <div class="chart-item">
                <div class="chart-label">{label}</div>
                <div class="chart-bar">
                    <div class="chart-fill" style="width: {rate}%; background: {color}"></div>
                </div>
                <div class="chart-value">{count} ({rate:.1f}%)</div>
            </div>
            '''
        
        html += '</div>'
        return html
    
    def _generate_detailed_results_html(self, fill_results: List[FillResult]) -> str:
        """生成详细结果表格的HTML"""
        if not fill_results:
            return '<div class="no-data">没有填充结果数据</div>'
        
        html = '''
        <table class="results-table">
            <thead>
                <tr>
                    <th>行</th>
                    <th>列</th>
                    <th>原内容</th>
                    <th>新内容</th>
                    <th>关系字段</th>
                    <th>填充源</th>
                    <th>置信度</th>
                    <th>状态</th>
                </tr>
            </thead>
            <tbody>
        '''
        
        for result in fill_results:
            # 处理内容长度
            original_content = result.original_content[:50] + ('...' if len(result.original_content) > 50 else '')
            new_content = result.new_content[:50] + ('...' if len(result.new_content) > 50 else '')
            
            # 处理关系字段
            relationships = ', '.join(result.relationship[:2]) if result.relationship else 'N/A'
            
            # 处理置信度
            if result.confidence > 0:
                confidence_class = 'confidence-high' if result.confidence > 0.9 else 'confidence-medium' if result.confidence >= 0.7 else 'confidence-low'
                confidence_text = f'<span class="{confidence_class}">{result.confidence:.2f}</span>'
            else:
                confidence_text = 'N/A'
            
            # 处理状态
            if result.success:
                status_text = '<span class="status-success">✓ 成功</span>'
            else:
                status_text = f'<span class="status-error">✗ 失败: {result.error_message}</span>'
            
            html += f'''
            <tr>
                <td>{result.row}</td>
                <td>{result.col}</td>
                <td>{original_content}</td>
                <td>{new_content}</td>
                <td>{relationships}</td>
                <td>{result.fill_source}</td>
                <td>{confidence_text}</td>
                <td>{status_text}</td>
            </tr>
            '''
        
        html += '''
            </tbody>
        </table>
        '''
        return html
    
    def _generate_error_details_html(self, fill_results: List[FillResult]) -> str:
        """生成错误详情的HTML"""
        error_results = [r for r in fill_results if not r.success]
        
        if not error_results:
            return ''
        
        html = '''
        <section class="error-section">
            <h2>❌ 错误详情</h2>
        '''
        
        for i, result in enumerate(error_results, 1):
            html += f'''
            <div class="error-item">
                <strong>错误 {i}:</strong> 行{result.row} 列{result.col} - {result.error_message}
            </div>
            '''
        
        html += '</section>'
        return html


class WordFormFiller:
    """Word表单填充器主类 - 协调各个组件完成填充任务"""
    
    def __init__(self, log_level: str = "INFO"):
        """初始化填充器"""
        self.setup_logging(log_level)
        self.logger = logging.getLogger(__name__)
        
        # 初始化组件
        self.fill_engine = FillEngine(log_level)
        self.report_generator = ReportGenerator()
    
    def setup_logging(self, level: str) -> None:
        """设置日志"""
        # 确保日志目录存在
        log_dir = '../../data/processed'
        os.makedirs(log_dir, exist_ok=True)
        
        logging.basicConfig(
            level=getattr(logging, level.upper()),
            format='%(asctime)s - %(name)s - %(levelname)s - %(message)s',
            handlers=[
                logging.FileHandler(os.path.join(log_dir, 'word_form_filler.log'), encoding='utf-8'),
                logging.StreamHandler()
            ]
        )
    
    def load_json_data(self, json_path: str) -> List[Dict[str, Any]]:
        """加载JSON数据"""
        return self.fill_engine.load_json_data(json_path)
    
    def fill_word_document(self, word_path: str, json_data: List[Dict[str, Any]], output_path: str = None) -> bool:
        """填充Word文档"""
        return self.fill_engine.fill_word_document(word_path, json_data, output_path)
    
    def generate_report(self, report_path: str = None, download_url: str = None) -> bool:
        """生成填充结果报告"""
        return self.report_generator.generate_report(self.fill_engine.fill_results, report_path, download_url)
    
    def get_statistics(self) -> Dict[str, Any]:
        """获取统计信息"""
        return self.fill_engine.get_statistics()


def create_cli_parser():
    """创建命令行参数解析器"""
    parser = argparse.ArgumentParser(
        description='Word表单填充工具 - 将JSON数据填写到Word文档中',
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="""
使用示例:
  # 基本填充
  python word_form_filler.py --json data.json --word template.docx
  
  # 指定输出路径
  python word_form_filler.py --json data.json --word template.docx --output filled.docx
  
  # 生成报告
  python word_form_filler.py --json data.json --word template.docx --report
  
  # 详细输出
  python word_form_filler.py --json data.json --word template.docx --verbose
        """
    )
    
    parser.add_argument(
        '--json', '-j',
        required=True,
        help='输入的JSON数据文件路径'
    )
    
    parser.add_argument(
        '--word', '-w',
        required=True,
        help='输入的Word模板文件路径 (.docx格式)'
    )
    
    parser.add_argument(
        '--output', '-o',
        help='输出的Word文件路径 (可选，默认自动生成)'
    )
    
    parser.add_argument(
        '--report', '-r',
        action='store_true',
        help='生成填充结果报告'
    )
    
    parser.add_argument(
        '--report-path',
        help='报告文件路径 (可选，默认自动生成)'
    )
    
    parser.add_argument(
        '--log-level',
        choices=['DEBUG', 'INFO', 'WARNING', 'ERROR'],
        default='INFO',
        help='日志级别 (默认: INFO)'
    )
    
    parser.add_argument(
        '--verbose', '-v',
        action='store_true',
        help='显示详细输出信息'
    )
    
    return parser


def main():
    """CLI主入口函数"""
    parser = create_cli_parser()
    args = parser.parse_args()
    
    try:
        # 验证输入文件
        if not os.path.exists(args.json):
            print(f"✗ 错误: JSON文件不存在: {args.json}")
            return 1
        
        if not os.path.exists(args.word):
            print(f"✗ 错误: Word文件不存在: {args.word}")
            return 1
        
        # 验证Word文件格式
        if not args.word.lower().endswith('.docx'):
            print(f"✗ 错误: 不支持的Word文件格式。请使用 .docx 文件")
            return 1
        
        # 设置日志级别
        log_level = "DEBUG" if args.verbose else args.log_level
        
        # 创建填充器
        filler = WordFormFiller(log_level=log_level)
        
        print(f"=== Word表单填充工具 ===")
        print(f"JSON数据文件: {args.json}")
        print(f"Word模板文件: {args.word}")
        print(f"日志级别: {log_level}")
        
        # 加载JSON数据
        print("\n📖 加载JSON数据...")
        json_data = filler.load_json_data(args.json)
        if not json_data:
            print("✗ 错误: 无法加载JSON数据")
            return 1
        
        print(f"✓ 成功加载 {len(json_data)} 条数据")
        
        # 生成输出文件路径
        if args.output:
            output_path = args.output
        else:
            word_file = Path(args.word)
            output_path = str(word_file.parent / f"{word_file.stem}_filled{word_file.suffix}")
        
        print(f"输出文件: {output_path}")
        
        # 执行填充
        print("\n🔄 开始填充...")
        success = filler.fill_word_document(args.word, json_data, output_path)
        
        if success:
            print(f"✓ 填充完成: {output_path}")
            
            # 生成报告
            if args.report:
                report_path = args.report_path
                if not report_path:
                    word_file = Path(args.word)
                    report_path = str(word_file.parent / f"{word_file.stem}_填充报告.html")
                
                print(f"\n📊 生成报告: {report_path}")
                if filler.generate_report(report_path):
                    print(f"✓ 报告生成完成: {report_path}")
                else:
                    print("✗ 报告生成失败")
            
            # 显示统计信息
            stats = filler.get_statistics()
            print(f"\n📈 填充统计:")
            print(f"  总数量: {stats['total_processed']}")
            print(f"  成功: {stats['successful_fills']}")
            print(f"  失败: {stats['failed_fills']}")
            print(f"  成功率: {stats['success_rate']:.1%}")
            
            if stats['average_confidence'] > 0:
                print(f"  平均置信度: {stats['average_confidence']:.2f}")
            
            # 按填充源统计
            if stats['source_statistics'] and args.verbose:
                print(f"\n📋 按填充源统计:")
                for source, data in stats['source_statistics'].items():
                    rate = data['success'] / data['total'] if data['total'] > 0 else 0
                    print(f"  {source}: {data['success']}/{data['total']} ({rate:.1%})")
            
            print(f"\n🎉 填充完成!")
            return 0
        
        else:
            print("✗ 填充失败")
            return 1
        
    except KeyboardInterrupt:
        print("\n✗ 用户中断操作")
        return 1
    except Exception as e:
        print(f"✗ 填充过程中发生错误: {e}")
        if args.verbose:
            import traceback
            traceback.print_exc()
        return 1


def demo_main():
    """演示主函数（保持向后兼容）"""
    # 配置参数
    JSON_PATH = r"D:\code\FF2\empty_cells_simplified_filled.json"
    WORD_PATH = r"D:\code\FF2\03d1def2c097db53b92c44eb974d26e4.docx"
    
    # 检查文件是否存在
    if not os.path.exists(JSON_PATH):
        print(f"错误: JSON文件不存在 - {JSON_PATH}")
        return
    
    if not os.path.exists(WORD_PATH):
        print(f"错误: Word文件不存在 - {WORD_PATH}")
        return
    
    # 创建填充器
    filler = WordFormFiller(log_level="INFO")
    
    print("开始Word表单填充...")
    print(f"JSON数据文件: {JSON_PATH}")
    print(f"Word模板文件: {WORD_PATH}")
    
    # 加载JSON数据
    json_data = filler.load_json_data(JSON_PATH)
    if not json_data:
        print("错误: 无法加载JSON数据")
        return
    
    # 生成输出文件路径
    word_file = Path(WORD_PATH)
    output_path = str(word_file.parent / f"{word_file.stem}_filled{word_file.suffix}")
    
    # 执行填充
    success = filler.fill_word_document(WORD_PATH, json_data, output_path)
    
    if success:
        print(f"✓ 填充完成，输出文件: {output_path}")
        
        # 生成报告
        report_path = str(word_file.parent / f"{word_file.stem}_填充报告.html")
        if filler.generate_report(report_path):
            print(f"✓ 报告生成完成: {report_path}")
        
        # 打印统计信息
        stats = filler.get_statistics()
        print("\n=== 填充统计 ===")
        print(f"总处理数量: {stats['total_processed']}")
        print(f"成功填充: {stats['successful_fills']}")
        print(f"失败数量: {stats['failed_fills']}")
        print(f"成功率: {stats['success_rate']:.1%}")
        print(f"平均置信度: {stats['average_confidence']:.2f}")
        
        print("\n按填充源统计:")
        for source, data in stats['source_statistics'].items():
            rate = data['success'] / data['total'] if data['total'] > 0 else 0
            print(f"  {source}: {data['success']}/{data['total']} ({rate:.1%})")
        
        print("\n置信度分布:")
        for level, count in stats['confidence_distribution'].items():
            print(f"  {level}: {count}")
        
    else:
        print("✗ 填充失败")


if __name__ == "__main__":
    # 如果没有命令行参数，运行演示模式
    import sys
    if len(sys.argv) == 1:
        demo_main()
    else:
        exit(main())